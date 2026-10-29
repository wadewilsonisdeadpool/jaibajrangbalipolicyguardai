"""Read, create and modify-a-copy for txt/md/csv/json/pdf/docx/xlsx.
Originals are opened read-only; edits always go to a new file."""
import csv
import io
import json
import shutil
from pathlib import Path
from typing import Any

MAX_TEXT_CHARS = 6000


def read_text(path: Path) -> tuple[str, bool]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        text = "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    elif suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=False)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True, max_row=300):
                if any(v is not None for v in row):
                    lines.append(",".join("" if v is None else str(v) for v in row))
        wb.close()
        text = "\n".join(lines)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    return text[:MAX_TEXT_CHARS], len(text) > MAX_TEXT_CHARS


def _parse_rows(content: str) -> list[list[Any]]:
    content = content.strip()
    if content.startswith("["):
        try:
            data = json.loads(content)
            if isinstance(data, list):
                return [row if isinstance(row, list) else [row] for row in data]
        except json.JSONDecodeError:
            pass
    return [row for row in csv.reader(io.StringIO(content))]


def _coerce(value: Any) -> Any:
    if isinstance(value, str):
        v = value.strip()
        try:
            return int(v) if v.lstrip("-").isdigit() else float(v)
        except ValueError:
            return value
    return value


def create(target: Path, content: str) -> None:
    suffix = target.suffix.lower()
    if suffix == ".docx":
        import docx
        d = docx.Document()
        for line in content.splitlines():
            s = line.rstrip()
            if s.startswith("### "):
                d.add_heading(s[4:], level=3)
            elif s.startswith("## "):
                d.add_heading(s[3:], level=2)
            elif s.startswith("# "):
                d.add_heading(s[2:], level=1)
            elif s.lstrip().startswith(("- ", "* ")):
                d.add_paragraph(s.lstrip()[2:], style="List Bullet")
            elif s:
                d.add_paragraph(s)
        d.save(str(target))
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for row in _parse_rows(content):
            ws.append([_coerce(v) for v in row])
        wb.save(str(target))
    elif suffix in {".md", ".txt", ".csv", ".json", ".html", ".log"}:
        target.write_text(content, encoding="utf-8")
    else:
        raise ValueError(f"Unsupported output type '{suffix}'. Use md, txt, csv, json, docx or xlsx.")


def modify_copy(source: Path, target: Path, args: dict[str, Any]) -> list[str]:
    """Copy source -> target, then apply edits to target only."""
    shutil.copy2(source, target)
    suffix = target.suffix.lower()
    changes: list[str] = []
    find, replace = args.get("find"), args.get("replace")
    append_text = args.get("append_text")

    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(target))
        for ref, value in (args.get("cell_updates") or {}).items():
            sheet_name, cell = (ref.split("!", 1) if "!" in ref else (wb.active.title, ref))
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Sheet '{sheet_name}' not found")
            wb[sheet_name][cell] = _coerce(value)
            changes.append(f"set {sheet_name}!{cell}")
        for row in args.get("append_rows") or []:
            wb.active.append([_coerce(v) for v in row])
            changes.append("appended a row")
        wb.save(str(target))
    elif suffix == ".docx":
        import docx
        d = docx.Document(str(target))
        if find:
            n = 0
            for p in d.paragraphs:
                if find in p.text:
                    p.text = p.text.replace(find, replace or "")
                    n += 1
            changes.append(f"replaced text in {n} paragraph(s)")
        if append_text:
            for line in str(append_text).splitlines():
                d.add_paragraph(line)
            changes.append("appended text")
        d.save(str(target))
    elif suffix == ".csv":
        rows = args.get("append_rows") or []
        if rows:
            with target.open("a", newline="", encoding="utf-8") as fh:
                csv.writer(fh).writerows(rows)
            changes.append(f"appended {len(rows)} row(s)")
        if find or append_text:
            text = target.read_text(encoding="utf-8")
            if find:
                text = text.replace(find, replace or "")
                changes.append("find/replace")
            if append_text:
                text += "\n" + str(append_text)
                changes.append("appended text")
            target.write_text(text, encoding="utf-8")
    elif suffix in {".txt", ".md", ".json", ".log"}:
        text = target.read_text(encoding="utf-8")
        if find:
            text = text.replace(find, replace or "")
            changes.append("find/replace")
        if append_text:
            text += "\n" + str(append_text)
            changes.append("appended text")
        target.write_text(text, encoding="utf-8")
    else:
        raise ValueError(f"Editing '{suffix}' files is not supported.")

    if not changes:
        changes.append("copied without changes (no edits given)")
    return changes
